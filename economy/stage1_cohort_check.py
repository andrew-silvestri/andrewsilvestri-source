"""Stage 1 feasibility check: does the product form predict performance in
any per-runner dataset we can actually download?

A. Lanferdini et al. 2020, Front Physiol 11:979 (CC BY), supplementary
   Table_1.xlsx: 20 recreational men, VO2max, VT1, VT2, economy at 12 and
   16 km/h, 3000 m time.  Joyner-form prediction: sustainable VO2 (taken as
   VT2, the second ventilatory threshold) divided by O2 cost per km at
   16 km/h gives a predicted speed.  Correlate with actual 3000 m speed and
   with each term alone.

B. Vickers & Vertosick 2016, BMC Sports Sci Med Rehabil 8:26 (CC BY 4.0),
   Additional file 2: 2,303 recreational runners with self-reported 5 km
   and marathon times.  Ratio of marathon speed to 5 km speed is a
   performance-level proxy for the sustainable fraction, and the Riegel
   exponent k in t2 = t1 (d2/d1)^k is re-fitted.

No design decisions.  Writes cohort_rows.json.
"""
import glob, json, math
import openpyxl

out = {}

# ---- A. Lanferdini ---------------------------------------------------------
p = glob.glob("sources/frontiers_2020*/*.xlsx")[0]
ws = openpyxl.load_workbook(p, read_only=True, data_only=True)[
    openpyxl.load_workbook(p, read_only=True).sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
hdr = rows[2]
data = [r for r in rows[3:] if r and isinstance(r[0], (int, float))]
col = {h: i for i, h in enumerate(hdr) if h}
vvo2 = [r[col['vVO2max (km/h)']] for r in data]
t3k = [r[col['3000m performance (s)']] for r in data]
vo2max = [r[col['VO2max (ml/kg/min)']] for r in data]
vt2 = [r[col['VT2 (ml/kg/min)']] for r in data]
re16 = [r[col['RE16 (ml/kg/min)']] for r in data]
n = len(data)
speed3k = [3.0 / (t / 3600) for t in t3k]           # km/h
cost = [re / 16.0 for re in re16]                    # ml/kg/min per km/h
frac = [v / m for v, m in zip(vt2, vo2max)]
pred = [v / c for v, c in zip(vt2, cost)]            # km/h, Joyner form
pred_max = [v / c for v, c in zip(vo2max, cost)]     # if fraction were 1

def pearson(x, y):
    mx, my = sum(x) / len(x), sum(y) / len(y)
    sxy = sum((a - mx) * (b - my) for a, b in zip(x, y))
    sxx = sum((a - mx) ** 2 for a in x); syy = sum((b - my) ** 2 for b in y)
    return sxy / math.sqrt(sxx * syy)

print(f"A. Lanferdini 2020, n={n}")
print(f"   3000 m time {min(t3k):.0f}-{max(t3k):.0f} s; VO2max {min(vo2max):.1f}-{max(vo2max):.1f};"
      f" VT2/VO2max {min(frac):.2f}-{max(frac):.2f}; RE16 {min(re16):.1f}-{max(re16):.1f} ml/kg/min"
      f" ({min(re16)/16*60:.0f}-{max(re16)/16*60:.0f} ml/kg/km)")
corr = {
    "VO2max": pearson(vo2max, speed3k),
    "VT2 (ml/kg/min)": pearson(vt2, speed3k),
    "fraction VT2/VO2max": pearson(frac, speed3k),
    "economy at 16 km/h (higher = worse)": pearson(re16, speed3k),
    "product VT2/cost": pearson(pred, speed3k),
    "product VO2max/cost": pearson(pred_max, speed3k),
    "vVO2max (measured)": pearson(vvo2, speed3k),
}
for k, v in corr.items():
    print(f"   r({k}, 3000 m speed) = {v:+.2f}")
ratio = [s / pm for s, pm in zip(speed3k, pred_max)]
print(f"   3000 m speed as share of VO2max/cost: mean {sum(ratio)/n:.2f}, range {min(ratio):.2f}-{max(ratio):.2f}")
out["lanferdini"] = dict(n=n, r=corr, share_mean=sum(ratio) / n)

# ---- B. Vickers & Vertosick -----------------------------------------------
p = glob.glob("sources/vickers*.xlsx")[0]
wb = openpyxl.load_workbook(p, read_only=True)
dd = list(wb[wb.sheetnames[1]].iter_rows(values_only=True))
print("\nB. Vickers & Vertosick 2016 data dictionary (time columns):")
for r in dd:
    if r and r[0] and any(str(r[0]).startswith(pfx) for pfx in ("k5_", "mf_", "endur")):
        print("   ", r[0], "=", r[1])
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
h = next(it); c = {k: i for i, k in enumerate(h)}
pairs = []
for r in it:
    t5, tm, d5, dm = r[c['k5_ti']], r[c['mf_ti']], r[c['k5_d']], r[c['mf_d']]
    if all(isinstance(x, (int, float)) and x for x in (t5, tm, d5, dm)):
        pairs.append((t5, tm, d5, dm))
ratios = sorted((tm / dm) ** -1 / (t5 / d5) ** -1 for t5, tm, d5, dm in pairs)   # marathon speed / 5k speed
ks = sorted(math.log(tm / t5) / math.log(dm / d5) for t5, tm, d5, dm in pairs)
q = lambda a, f: a[int(f * (len(a) - 1))]
print(f"   runners with both 5 km and marathon: {len(pairs)}")
print(f"   marathon speed / 5 km speed: median {q(ratios,.5):.3f}, IQR {q(ratios,.25):.3f}-{q(ratios,.75):.3f}, 5-95% {q(ratios,.05):.3f}-{q(ratios,.95):.3f}")
print(f"   Riegel exponent k: median {q(ks,.5):.3f}, IQR {q(ks,.25):.3f}-{q(ks,.75):.3f} (Riegel 1981 quoted 1.06-1.08)")
# does the ratio depend on ability?  bin by 5 km speed tercile
byspeed = sorted(((d5 / t5), (tm / dm) ** -1 / (t5 / d5) ** -1) for t5, tm, d5, dm in pairs)
terc = len(byspeed) // 3
for i, name in enumerate(("slowest third", "middle third", "fastest third")):
    seg = byspeed[i * terc:(i + 1) * terc]
    print(f"   {name}: median ratio {sorted(s[1] for s in seg)[len(seg)//2]:.3f}")
out["vickers"] = dict(n=len(pairs), ratio_median=q(ratios, .5), k_median=q(ks, .5),
                      ratio_iqr=[q(ratios, .25), q(ratios, .75)])
json.dump(out, open("cohort_rows.json", "w"), indent=1)
